from pathlib import Path
import hashlib
import openpyxl

REQUIRED_COLUMNS = [
    "item_id", "concept_id", "language", "language_code", "task_code",
    "task_family", "item_no", "dataset_role", "target_script",
    "English prompt", "English reference answer",
    "FINAL target-language prompt", "FINAL target-language reference answer",
    "expert/validation decision", "final_status", "validation_provenance"
]

EXPECTED_LANGS = {"SAN","SAT","BOD","DOI","KON","KAS","MAI","MNI","SND"}
EXPECTED_TASKS = {"FK","EX","CC","SM","SP","EU"}

def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()

def load_and_validate(path, expected_sha256=None):
    path = Path(path)
    if expected_sha256 and sha256_file(path) != expected_sha256:
        raise ValueError("Benchmark SHA-256 mismatch: file is not the locked benchmark.")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Master_1080" not in wb.sheetnames:
        raise ValueError("Missing Master_1080 sheet.")

    ws = wb["Master_1080"]
    rows = list(ws.iter_rows(values_only=True))
    headers = list(rows[0])
    missing_cols = [c for c in REQUIRED_COLUMNS if c not in headers]
    if missing_cols:
        raise ValueError(f"Missing required columns: {missing_cols}")

    records = [dict(zip(headers, row)) for row in rows[1:] if any(v is not None for v in row)]

    assert len(records) == 1080, f"Expected 1080 items; found {len(records)}"
    ids = [r["item_id"] for r in records]
    assert len(ids) == len(set(ids)), "Duplicate item_id detected"
    assert {r["language_code"] for r in records} == EXPECTED_LANGS
    assert {r["task_code"] for r in records} == EXPECTED_TASKS

    for code in EXPECTED_LANGS:
        subset = [r for r in records if r["language_code"] == code]
        assert len(subset) == 120, f"{code}: expected 120, found {len(subset)}"
        assert sum(r["dataset_role"] == "Parallel" for r in subset) == 90
        assert sum(r["dataset_role"] == "Contextual" for r in subset) == 30
        for task in EXPECTED_TASKS:
            assert sum(r["task_code"] == task for r in subset) == 20

    required_text = [
        "English prompt", "English reference answer",
        "FINAL target-language prompt", "FINAL target-language reference answer"
    ]
    for r in records:
        for col in required_text:
            if r[col] is None or not str(r[col]).strip():
                raise ValueError(f"Missing {col} in {r['item_id']}")

    return records
